from  docx import Document 
from openpyxl import Workbook

def generate_word_report(data):
    doc = Document()
    doc.add_heading('Отчёт по почтовым отправлениям', level=1)

    for item in data:
        doc.add_paragrapg(f"Вес:{item['weight']}кг, Расстояние; {item['distance']} км, Цена; {item['cost']} руб.")

    doc.save('postage_report.docx')
def generate_excel_report(data):
    wb = Workbook()
    wb = wb.active
    ws.append(["Вес","Расстояние", "Цена"])

    for item in data:
        ws.append([item['weight'], item['distance'], item['cost']])

    wb.save('postage_report.xls')